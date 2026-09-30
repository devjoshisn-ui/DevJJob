#!/usr/bin/env python3
"""
Gera a versão Excel (.xlsx) do relatório: links clicáveis, colunas largas,
cabeçalho congelado, filtros e cores por recomendação. O screen.py já chama
isto no fim de cada rodada; rode à mão só pra converter um CSV antigo:

    python exportar_excel.py                                   # CSV mais recente em output/
    python exportar_excel.py output/vagas_triadas_2026-09-21.csv
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# (campo do CSV, cabeçalho na planilha, largura)
COLUNAS = [
    ("recomendacao", "Recomendação", 16),
    ("nota_fit", "Nota", 7),
    ("titulo", "Vaga", 60),
    ("empresa", "Empresa", 28),
    ("link", "Link", 12),
    ("trilha", "Trilha", 24),
    ("cidade", "Cidade", 22),
    ("remoto", "Remoto", 9),
    ("salario", "Salário (Gupy / remoto internacional)", 30),
    ("senioridade_ok", "Senioridade ok", 14),
    ("gap_critico", "Gap crítico", 12),
    ("aderencia_nivel", "Aderência", 50),
    ("ajuste", "Ajuste", 45),
    ("confianca", "Confiança", 11),
    ("fonte", "Fonte", 13),
    ("publicado_em", "Publicada em", 22),
]

CORES = {
    "aplicar_agora": "C6EFCE",  # verde
    "avaliar": "FFEB9C",  # amarelo
    "baixa_prioridade": "F4CCCC",  # rosa
    "nao_aplicar": "E7E6E6",  # cinza
}

FONTE = "Arial"


def write_xlsx(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Vagas"

    for c, (_, cab, larg) in enumerate(COLUNAS, 1):
        cell = ws.cell(row=1, column=c, value=cab)
        cell.font = Font(name=FONTE, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = larg

    for i, r in enumerate(rows, 2):
        fill = PatternFill("solid", fgColor=CORES[r["recomendacao"]]) if r.get("recomendacao") in CORES else None
        for c, (campo, _, _) in enumerate(COLUNAS, 1):
            valor = r.get(campo, "")
            cell = ws.cell(row=i, column=c)
            cell.font = Font(name=FONTE)
            if campo == "link" and valor:
                cell.value = "Abrir vaga"
                cell.hyperlink = valor
                cell.font = Font(name=FONTE, color="0563C1", underline="single")
            elif campo in ("nota_fit", "confianca") and valor not in ("", None):
                cell.value = float(valor)
                cell.number_format = "0.0" if campo == "nota_fit" else "0.00"
            elif campo == "remoto":
                cell.value = "sim" if str(valor).lower() in ("true", "1", "sim") else "não"
            else:
                cell.value = str(valor).strip() if valor is not None else ""
            if fill and campo in ("recomendacao", "nota_fit"):
                cell.fill = fill

    ws.freeze_panes = "D2"  # cabeçalho + recomendação/nota/vaga sempre visíveis
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{max(len(rows) + 1, 1)}"
    wb.save(path)


def main():
    if len(sys.argv) > 1:
        csv_path = Path(sys.argv[1])
    else:
        candidatos = sorted(Path("output").glob("vagas_triadas_*.csv"))
        if not candidatos:
            sys.exit("Nenhum output/vagas_triadas_*.csv encontrado.")
        csv_path = candidatos[-1]
    with csv_path.open(encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    out = csv_path.with_suffix(".xlsx")
    write_xlsx(rows, out)
    print(f"Planilha salva em: {out}")


if __name__ == "__main__":
    main()
